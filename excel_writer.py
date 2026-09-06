from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from invoice_reader import InvoiceRecord
from reimbursement_rules import (
    REIMBURSEMENT_TYPES,
    ReimbursementRule,
    classify_reimbursement,
    normalize_reimbursement_type,
)


HEADERS = [
    "实报人",
    "学号",
    "发票上的！发票上的！发票上的！不要瞎写！项目名称",
    "发票上的！发票上的！发票上的！不要瞎写！规格型号",
    "报销类型",
    "票面金额",
    "填报日期",
    "电话",
    "是否传递发票",
    "备注",
]
MANUAL_REVIEW_NOTE = "已人工核对票面项目名称及规格型号"


@dataclass(frozen=True)
class ReimbursementProfile:
    claimant: str = "李府鸿"
    student_id: str = "12604040"
    reimbursement_type: str = "材料费"
    phone: str = "13516389370"
    transmit_invoice: str = "YES"


def build_workbook(
    records: list[InvoiceRecord],
    fill_date: datetime,
    profile: ReimbursementProfile,
    rules: tuple[ReimbursementRule, ...] | None = None,
    confirmed_review_ids: set[str] | None = None,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "发票报销填报表"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    fallback_type = normalize_reimbursement_type(profile.reimbursement_type)
    classifications = [
        classify_reimbursement(record.item_name, fallback_type, rules)
        for record in records
    ]

    confirmed_ids = confirmed_review_ids or set()
    sheet.append(HEADERS)
    for index, (record, classification) in enumerate(
        zip(records, classifications, strict=True),
        start=1,
    ):
        remarks = list(classification.remarks)
        if f"invoice-{index:03d}" in confirmed_ids:
            remarks.append(MANUAL_REVIEW_NOTE)
        remark_text = "\n".join(remarks)
        sheet.append([
            profile.claimant,
            profile.student_id,
            record.item_name,
            record.specification,
            classification.reimbursement_type,
            float(record.amount),
            fill_date,
            profile.phone,
            profile.transmit_invoice,
            remark_text,
        ])

    total_row = len(records) + 2
    sheet.cell(total_row, 5, "票面金额合计")
    sheet.cell(total_row, 6, float(sum(record.amount for record in records)))

    header_font = Font(name="Arial", size=10, bold=True, color="000000")
    body_font = Font(name="Arial", size=10, color="000000")
    remark_font = Font(name="Arial", size=10, color="9C0006")
    thin_black = Side(style="thin", color="000000")
    solid_border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
    no_fill = PatternFill(fill_type=None)
    warning_fill = PatternFill("solid", fgColor="FFC7CE")

    for cell in sheet[1]:
        cell.fill = no_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = solid_border
    sheet.row_dimensions[1].height = 52

    for row in sheet.iter_rows(min_row=2, max_row=total_row, min_col=1, max_col=10):
        for cell in row:
            cell.fill = no_fill
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = solid_border

    for row_number, classification in enumerate(classifications, start=2):
        for column in (1, 2, 3, 4, 5, 8, 9, 10):
            sheet.cell(row_number, column).data_type = "s"
        sheet.cell(row_number, 2).number_format = "@"
        sheet.cell(row_number, 6).number_format = "0.00"
        sheet.cell(row_number, 7).number_format = "yyyy/m/d"
        sheet.cell(row_number, 8).number_format = "@"
        sheet.cell(row_number, 6).alignment = Alignment(horizontal="right", vertical="top")
        sheet.cell(row_number, 7).alignment = Alignment(horizontal="center", vertical="top")
        remark_cell = sheet.cell(row_number, 10)
        if remark_cell.value:
            remark_cell.fill = warning_fill
            remark_cell.font = remark_font
            first_url = next((url for url in classification.template_urls if url), None)
            if first_url:
                remark_cell.hyperlink = first_url
                remark_cell.font = Font(name="Arial", size=10, color="9C0006", underline="single")

    for cell in sheet[total_row]:
        cell.fill = no_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="000000")
    sheet.cell(total_row, 6).number_format = "0.00"

    widths = [11, 14, 48, 32, 16, 14, 14, 18, 16, 72]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_number in range(2, len(records) + 2):
        line_count = max(1, str(sheet.cell(row_number, 10).value or "").count("\n") + 1)
        sheet.row_dimensions[row_number].height = max(45, 30 * line_count)

    sheet.auto_filter.ref = f"A1:J{len(records) + 1}"
    if records:
        reimbursement_validation = DataValidation(
            type="list",
            formula1='"' + ",".join(REIMBURSEMENT_TYPES) + '"',
            allow_blank=False,
        )
        reimbursement_validation.error = "请从规定的报销类型中选择。"
        reimbursement_validation.errorTitle = "报销类型无效"
        reimbursement_validation.showErrorMessage = True
        sheet.add_data_validation(reimbursement_validation)
        reimbursement_validation.add(f"E2:E{len(records) + 1}")

        transmit_validation = DataValidation(type="list", formula1='"YES,NO"', allow_blank=False)
        transmit_validation.error = "只能选择 YES 或 NO。"
        transmit_validation.errorTitle = "是否传递发票无效"
        transmit_validation.showErrorMessage = True
        sheet.add_data_validation(transmit_validation)
        transmit_validation.add(f"I2:I{len(records) + 1}")
    return workbook


def build_excel_bytes(
    records: list[InvoiceRecord],
    fill_date: datetime,
    profile: ReimbursementProfile,
    rules: tuple[ReimbursementRule, ...] | None = None,
    confirmed_review_ids: set[str] | None = None,
) -> bytes:
    workbook = build_workbook(records, fill_date, profile, rules, confirmed_review_ids)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def write_excel(
    records: list[InvoiceRecord],
    output_path: Path,
    fill_date: datetime,
    profile: ReimbursementProfile | None = None,
    rules: tuple[ReimbursementRule, ...] | None = None,
    confirmed_review_ids: set[str] | None = None,
) -> None:
    content = build_excel_bytes(
        records,
        fill_date,
        profile or ReimbursementProfile(),
        rules,
        confirmed_review_ids,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix="invoice_reimbursement_",
        suffix=".xlsx",
        dir=output_path.parent,
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        temporary_path.write_bytes(content)
        os.replace(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)
