"""隔离检查：只使用临时数据库，不写入现有配置或实际报销历史。"""
from __future__ import annotations

import io
import json
import os
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path


def main() -> None:
    root = Path(__file__).resolve().parent
    temporary_root = root / "tmp"
    temporary_root.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(dir=temporary_root, prefix="workflow_check_") as temporary:
        directory = Path(temporary)
        os.environ["RULES_DB_PATH"] = str(directory / "check.db")
        os.environ["SITE_ASSET_DIR"] = str(directory / "assets")

        from fastapi.testclient import TestClient
        from openpyxl import load_workbook
        from pypdf import PdfReader
        from reportlab.pdfgen import canvas

        from invoice_history import find_history, list_history, save_history
        from invoice_review import apply_review, prepare_explanations
        from processing import process_inputs
        from reimbursement_rules import get_rules
        from server import app
        from site_settings import get_site_texts, update_site_texts

        document = io.BytesIO()
        page = canvas.Canvas(document, pagesize=(595, 300))
        page.drawString(50, 200, "FIRST-INVOICE")
        page.showPage()
        page.drawString(50, 200, "SECOND-INVOICE")
        page.save()
        source = directory / "merged.pdf"
        source.write_bytes(document.getvalue())
        batch = process_inputs([source])
        assert len(batch.records) == 2
        broken = directory / "broken.pdf"
        broken.write_bytes(b"%PDF-unreadable")
        damaged_batch = process_inputs([broken, source])
        assert len(damaged_batch.records) == 3 and damaged_batch.records[0].needs_review
        # 每次重读同一合并 PDF，页指纹必须保持一致。
        assert [r.source_digest for r in batch.records] == [r.source_digest for r in process_inputs([source]).records]
        edits = [{
            "record_id": f"invoice-{index + 1:03d}",
            "pdf_name": record.pdf_name,
            "source_digest": record.source_digest,
            "item_name": "*电线电缆*连接线",
            "specification": "",
            "invoice_number": f"9900000000000000000{index}",
            "amount": "19.49",
            "confirmed": True,
            "excluded": index == 1,
        } for index, record in enumerate(batch.records)]
        corrected, indexes, confirmed = apply_review(batch, json.dumps(edits), set())
        assert indexes == [0] and confirmed == {"invoice-001"}
        assert corrected.validation.can_export
        assert corrected.records[0].specification == ""
        unconfirmed = [dict(edits[0], confirmed=False), edits[1]]
        assert not apply_review(batch, json.dumps(unconfirmed), set())[0].validation.can_export
        for invalid in ("NaN", "Infinity", "-1", "19.491"):
            bad = [dict(edits[0], amount=invalid), edits[1]]
            try:
                apply_review(batch, json.dumps(bad), set())
            except ValueError:
                pass
            else:
                raise AssertionError(f"invalid amount accepted: {invalid}")

        notes = [{"rule_id": "cable", "note": "上述线材用于实验设备电源连接，特此说明。"}]
        prepared = prepare_explanations(corrected.records, "材料费", get_rules(), json.dumps(notes), "检查用户", "2026/10/20")
        assert len(prepared) == 1 and "19.49" in prepared[0]["details"]

        with TestClient(app) as client:
            form = {"edits": json.dumps(edits), "reimbursement_type": "材料费",
                    "claimant": "检查用户", "student_id": "CHECK-ONLY", "phone": "CHECK-PHONE",
                    "fill_date": "2026-10-20", "explanations": json.dumps(notes)}
            def post(url, fields=None):
                return client.post(url, data=fields or form, files={"files": ("merged.pdf", document.getvalue(), "application/pdf")})
            preview = post("/v1/invoices/preview")
            assert preview.status_code == 200 and len(preview.json()["records"]) == 2, preview.text
            exported = post("/v1/invoices/export")
            assert exported.status_code == 200, exported.text
            workbook = load_workbook(io.BytesIO(exported.content))
            assert workbook.active.max_row == 3 and workbook.active.cell(2, 6).value == 19.49
            workbook.close()
            printed = post("/v1/invoices/print")
            assert printed.status_code == 200, printed.text
            printable = PdfReader(io.BytesIO(printed.content))
            assert len(printable.pages) == 2
            assert "FIRST-INVOICE" in printable.pages[0].extract_text()
            assert "SECOND-INVOICE" not in printable.pages[0].extract_text()
            assert float(printable.pages[1].mediabox.height) > 840
            # 两张发票同页，说明页始终保持独立 A4。
            other_edits = [edits[0], dict(edits[1], excluded=False)]
            both = post("/v1/invoices/print", dict(form, edits=json.dumps(other_edits)))
            assert both.status_code == 200, both.text
            combined = PdfReader(io.BytesIO(both.content))
            assert len(combined.pages) == 2
            assert "SECOND-INVOICE" in combined.pages[0].extract_text()
            registered = post("/v1/invoices/history/register")
            assert registered.status_code == 200, registered.text
            assert post("/v1/invoices/export").status_code == 422
            assert len(list_history("CHECK-ONLY", "CHECK-PHONE")) == 1
            identity = {"student_id": "CHECK-ONLY", "phone": "CHECK-PHONE"}
            assert client.post("/v1/invoices/history/query", json=identity).status_code == 200
            assert client.post("/v1/invoices/history/undo", json=dict(identity, invoice_number=edits[0]["invoice_number"])).status_code == 200
            assert post("/v1/invoices/export").status_code == 200

        texts = get_site_texts()
        update_site_texts(dict(texts, hero_title="保留我的原配置"))
        get_rules()
        find_history([edits[0]["invoice_number"]])
        assert get_site_texts()["hero_title"] == "保留我的原配置"
        record = corrected.records[0]
        with ThreadPoolExecutor(max_workers=5) as pool:
            registrations = list(pool.map(lambda _: save_history([record], "检查用户", "CHECK-ONLY", "CHECK-PHONE"), range(5)))
        assert sum(not duplicates for duplicates in registrations) == 1

        print("PASS: manual edits, exclusions, Excel, two-up printing, full-page notes, history and concurrent duplicates, preserved settings")


if __name__ == "__main__":
    main()
